"""Validate document structure and the declared route model, without AWS access."""

import hashlib
import ipaddress
import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

import content_identity  # noqa: F401
import content_network
import content_operations  # noqa: F401
import content_reference  # noqa: F401
import content_diagrams  # noqa: F401
from report_model import SECTIONS
from report_sources import SOURCES

ROOT = Path(__file__).resolve().parents[1]
FILES = list((ROOT / "output").glob("*.xlsx"))
assert len(FILES) == 1, "Expected exactly one deliverable workbook"
XLSX = FILES[0]
MD = XLSX.with_suffix(".md")
text = MD.read_text(encoding="utf-8")
checks = []


def check(ok, message):
    assert ok, message
    checks.append(message)


with zipfile.ZipFile(XLSX) as z:
    check(z.testzip() is None, "XLSX ZIP integrity")
wb = load_workbook(XLSX, data_only=False)
check(len(wb.worksheets) == 27, "27 worksheets present")
check(len({s.title for s in wb}) == len(wb.worksheets), "Unique worksheet names")
check(wb.active.title == "00_読む順番", "Contents is initial worksheet")

nonempty = 0
for ws in wb:
    check(bool(ws.print_area), f"Print area: {ws.title}")
    check(bool(ws.freeze_panes), f"Freeze panes: {ws.title}")
    check(ws.sheet_view.showGridLines is False, f"Gridlines hidden: {ws.title}")
    for row in ws:
        for c in row:
            if c.value is None:
                continue
            nonempty += 1
            assert c.font.name == "Meiryo UI", (ws.title, c.coordinate, c.font.name)
            assert c.alignment.wrap_text, (ws.title, c.coordinate, "missing wrap")
            assert c.data_type != "f", (ws.title, c.coordinate, "unexpected formula")
            assert c.data_type != "e", (ws.title, c.coordinate, "cell error")
            assert "\ufffd" not in str(c.value), (ws.title, c.coordinate, "replacement character")
            assert ws.row_dimensions[c.row].height and ws.row_dimensions[c.row].height <= 409.5
            for color in [c.font.color, c.fill.fgColor if c.fill.patternType == "solid" else None]:
                if color is not None and color.type == "rgb":
                    rgb = str(color.rgb)[-6:]
                    assert rgb[:2] == rgb[2:4] == rgb[4:6], (ws.title, c.coordinate, rgb)
            if c.hyperlink and c.hyperlink.target and c.hyperlink.target.startswith("#"):
                target = c.hyperlink.target
                matched = re.match(r"#'(.+)'!([A-Z]+\d+)$", target)
                assert matched and matched[1] in wb.sheetnames, target
                assert wb[matched[1]][matched[2]].value is not None, target
check(True, f"Meiryo UI / wrapped / grayscale / valid hyperlinks: {nonempty} populated cells")

for sec in SECTIONS:
    check(sec["sheet"] in wb.sheetnames, f"Chapter: {sec['sheet']}")
    for i, row in enumerate(sec["rows"], 1):
        for reference in re.findall(r"S\d{2}", row[4]):
            assert reference in SOURCES, reference
        assert f"### {sec['number']:02d}.{i:02d} {row[0]}" in text
        if sec["kind"] != "architecture":
            for col, value in enumerate(row, 2):
                assert wb[sec["sheet"]].cell(i + 6, col).value == value
check(True, "All chapter entries agree between source data, workbook and Markdown")
check("\ufffd" not in text and not text.startswith("\ufeff"), "Markdown is UTF-8 without BOM or replacement characters")
check(text.count("```mermaid") == 2, "Two self-contained Mermaid diagrams")
check(text.count("```") % 2 == 0, "Balanced Markdown fences")
for key, source in SOURCES.items():
    assert f"### {key} " in text
    assert source["url"] in text
check(True, "All 68 primary-source links included")

route_tables = {}
for table, cidr, target, purpose in content_network.ROUTES:
    route_tables.setdefault(table, []).append((ipaddress.ip_network(cidr), target))


def route(table, dest):
    address = ipaddress.ip_address(dest)
    matches = [(net.prefixlen, target) for net, target in route_tables[table] if address in net]
    assert matches, (table, dest)
    return max(matches)[1]


route_results = []
for link_az, source in [("A", "10.20.10.10"), ("B", "10.20.11.10")]:
    for alb_az, dest in [("A", "10.20.30.10"), ("B", "10.20.31.10")]:
        forward = route(f"RT-Link-{link_az}", dest)
        reverse = route(f"RT-ALB-{alb_az}", source)
        assert forward == reverse == f"GWLBE-{link_az}"
        assert route(f"RT-Endpoint-{link_az}", source) == "local"
        assert route(f"RT-Endpoint-{link_az}", dest) == "local"
        route_results.append({"source": source, "destination": dest, "forward": forward, "reverse": reverse, "post_inspection": "local"})
check(True, "Four AZ-pair route examples are symmetric and exit inspection with local routes")
for task_az, task_ip in [("A", "10.20.40.50"), ("B", "10.20.41.50")]:
    for alb_az, alb_ip in [("A", "10.20.30.10"), ("B", "10.20.31.10")]:
        assert route(f"RT-ALB-{alb_az}", task_ip) == "local"
        assert route(f"RT-Task-{task_az}", alb_ip) == "local"
check(True, "ALB-to-ECS connections remain distinct local routes in the stated design")

# Detect whitespace problems in newly authored text files, without pretending this is a Git repository.
whitespace = []
for path in list((ROOT / "scripts").glob("*.py")) + list((ROOT / "scripts").glob("*.ps1")):
    data = path.read_text(encoding="utf-8-sig")
    for n, line in enumerate(data.splitlines(), 1):
        if line.rstrip() != line:
            whitespace.append((path.name, n))
    assert data.endswith("\n"), path
check(not whitespace, "Authored scripts have no trailing whitespace and have final newlines")

report = {
    "status": "PASS",
    "scope": "Local document and static routing-model validation; no live AWS or PAN-OS tests",
    "sheets": len(wb.worksheets),
    "entries": sum(len(s["rows"]) for s in SECTIONS),
    "sources": len(SOURCES),
    "populated_cells": nonempty,
    "checks": checks,
    "route_cases": route_results,
    "files": {p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in [XLSX, MD]},
    "git_diff_check": "Not applicable: workspace is not a Git repository; whitespace and content reviewed separately",
}
validation = ROOT / "validation"
validation.mkdir(exist_ok=True)
(validation / "static_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({k: report[k] for k in ["status", "sheets", "entries", "sources", "populated_cells"]}, ensure_ascii=True))

