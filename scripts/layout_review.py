"""Conservative cell-layout review with the installed Meiryo UI font.

This is a local approximation, explicitly not a Microsoft Excel screen capture.
It measures at narrow 7-pixel column units to avoid optimistic fitting estimates.
"""

import json
import math
import re
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
VALIDATION = ROOT / "validation"
VALIDATION.mkdir(exist_ok=True)
FILE = next((ROOT / "output").glob("*.xlsx"))
wb = load_workbook(FILE)


@lru_cache(maxsize=64)
def measured_font(size, bold):
    filename = "meiryob.ttc" if bold else "meiryo.ttc"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / filename), max(8, round(size * 4 / 3)), index=2)


def wrap_line(line, font, width):
    if not line:
        return [""]
    output = []
    current = ""
    tokens = re.findall(r"[A-Za-z0-9._:/%{}<>@+\-]+|[ \t]+|.", line)
    for token in tokens:
        if font.getlength(current + token) <= width:
            current += token
            continue
        if current:
            output.append(current.rstrip())
            current = ""
        if token.isspace():
            continue
        if font.getlength(token) <= width:
            current = token
        else:
            for ch in token:
                if current and font.getlength(current + ch) > width:
                    output.append(current)
                    current = ""
                current += ch
    if current or not output:
        output.append(current.rstrip())
    return output


def lines_for(value, font, width):
    return [part for line in str(value).split("\n") for part in wrap_line(line, font, width)]


def geometry(ws):
    xs = [0]
    ys = [0]
    for col in range(1, ws.max_column + 1):
        xs.append(xs[-1] + round(ws.column_dimensions[get_column_letter(col)].width * 7 + 5))
    for row in range(1, ws.max_row + 1):
        ys.append(ys[-1] + round((ws.row_dimensions[row].height or 15) * 4 / 3))
    merges = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    return xs, ys, merges


adjustments = []
for ws in wb:
    xs, ys, merges = geometry(ws)
    if ws.title.startswith("02_"):
        continue
    for r in range(7, ws.max_row + 1):
        need = 0
        for cell in ws[r]:
            if cell.value is None or (r, cell.column) in merges:
                continue
            font = measured_font(cell.font.sz or 11, bool(cell.font.bold))
            height = sum(font.getmetrics())
            width = xs[cell.column] - xs[cell.column - 1] - 12
            need = max(need, len(lines_for(cell.value, font, width)) * height + 10)
        old = ws.row_dimensions[r].height or 15
        proposed = math.ceil(need * 3 / 4) + 2
        if proposed > old:
            assert proposed <= 409.5, (ws.title, r, proposed)
            ws.row_dimensions[r].height = proposed
            adjustments.append({"sheet": ws.title, "row": r, "old_points": old, "new_points": proposed})

overflow = []
populated = 0
for ws in wb:
    xs, ys, merges = geometry(ws)
    for row in ws:
        for cell in row:
            if cell.value is None:
                continue
            populated += 1
            m = merges.get((cell.row, cell.column))
            max_col = m.max_col if m else cell.column
            max_row = m.max_row if m else cell.row
            width = xs[max_col] - xs[cell.column - 1] - 12
            height = ys[max_row] - ys[cell.row - 1]
            font = measured_font(cell.font.sz or 11, bool(cell.font.bold))
            required = len(lines_for(cell.value, font, width)) * sum(font.getmetrics()) + 8
            if required > height:
                overflow.append({"sheet": ws.title, "cell": cell.coordinate, "available_pixels": height, "needed_pixels": required})

# The generated table data are unchanged; only conservative row-height padding may differ.
wb.save(FILE)


def rgb(color, default):
    if color is not None and color.type == "rgb":
        return "#" + str(color.rgb)[-6:]
    return default


def render(ws, last_row, filename):
    xs, ys, merges = geometry(ws)
    last_row = min(last_row, ws.max_row)
    w = xs[-1]
    margin = 16
    label_height = 34
    img = Image.new("RGB", (w + margin * 2, ys[last_row] + margin * 2 + label_height), "#dddddd")
    draw = ImageDraw.Draw(img)
    draw.text((margin, 7), "セル寸法に基づくローカルプレビュー（Excelの実画面ではありません）", font=measured_font(10, False), fill="#303030")
    yoff = margin + label_height
    for row in ws.iter_rows(min_row=1, max_row=last_row):
        for cell in row:
            if cell.value is None:
                continue
            m = merges.get((cell.row, cell.column))
            end_col = m.max_col if m else cell.column
            end_row = m.max_row if m else cell.row
            if end_row > last_row:
                continue
            x1, y1 = xs[cell.column - 1] + margin, ys[cell.row - 1] + yoff
            x2, y2 = xs[end_col] + margin, ys[end_row] + yoff
            fill = rgb(cell.fill.fgColor, "#ffffff") if cell.fill.patternType == "solid" else "#ffffff"
            draw.rectangle((x1, y1, x2, y2), fill=fill)
            font = measured_font(cell.font.sz or 11, bool(cell.font.bold))
            lines = lines_for(cell.value, font, x2 - x1 - 12)
            line_height = sum(font.getmetrics())
            text_height = len(lines) * line_height
            top = y1 + 4
            if cell.alignment.vertical == "center":
                top = y1 + (y2 - y1 - text_height) / 2
            for line in lines:
                left = x1 + 6
                if cell.alignment.horizontal == "center":
                    left = x1 + ((x2 - x1) - font.getlength(line)) / 2
                elif cell.alignment.horizontal == "right":
                    left = x2 - 6 - font.getlength(line)
                draw.text((left, top), line, font=font, fill=rgb(cell.font.color, "#202020"), anchor="lt")
                top += line_height
            if cell.border.bottom.style:
                draw.line((x1, y2 - 1, x2, y2 - 1), fill="#c6c6c6", width=1)
    img.save(VALIDATION / filename)


render(wb["00_読む順番"], 16, "preview_contents.png")
render(wb["02_全体構成図"], 62, "preview_architecture.png")
render(wb["14_ルート表"], 13, "preview_routes.png")
render(wb["16_TLSと検査範囲"], 12, "preview_tls.png")
render(wb["26_HTTP例と設定対応"], 11, "preview_http.png")

result = {
    "status": "PASS" if not overflow else "NEEDS_REVIEW",
    "method": "Conservative local approximation using installed Meiryo UI, measured wrapped text, and workbook dimensions",
    "native_excel_rendering": False,
    "populated_cells_checked": populated,
    "height_adjustments": adjustments,
    "overflow": overflow,
    "previews": ["preview_contents.png", "preview_architecture.png", "preview_routes.png", "preview_tls.png", "preview_http.png"],
}
(VALIDATION / "layout_validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": result["status"], "checked_cells": populated, "adjusted_rows": len(adjustments), "overflow": overflow}, ensure_ascii=False))
assert not overflow, "Fix overflow before delivery"

