"""Build Japanese, monochrome Excel and standalone Markdown deliverables."""

from __future__ import annotations

import json
import math
import re
from datetime import datetime
from pathlib import Path

from PIL import ImageFont
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

import content_identity  # noqa: F401
import content_network  # noqa: F401
import content_operations  # noqa: F401
import content_reference  # noqa: F401
import content_diagrams
from report_model import SECTIONS
from report_sources import AS_OF, SOURCES

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
STEM = "AWS_M2M_CloudFront_GWLBE_詳細解説_20261003"
XLSX = OUT / f"{STEM}.xlsx"
MD = OUT / f"{STEM}.md"
FONT = "Meiryo UI"
DARK = "242424"
INK = "202020"
MID = "606060"
LIGHT = "F1F1F1"
WHITE = "FFFFFF"
LINE = "C6C6C6"
SECTIONS.sort(key=lambda s: s["number"])
SOURCE_ROWS = {key: i + 7 for i, key in enumerate(SOURCES)}
BODY_WIDTHS = [7, 25, 39, 70, 49, 15]
MEASURE = ImageFont.truetype(r"C:\Windows\Fonts\meiryo.ttc", 15, index=2)


def font(size=11, bold=False, color=INK, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline)


def wrapped_lines(value, pixel_width):
    lines = 0
    for paragraph in str(value).split("\n"):
        if not paragraph:
            lines += 1
            continue
        used = 0.0
        lines += 1
        for char in paragraph:
            width = MEASURE.getlength(char)
            if used + width > pixel_width and used:
                lines += 1
                used = 0.0
            used += width
    return lines


def row_height(values, widths, minimum=50):
    # Conservative initial height. Excel AutoFit is applied in a separate validation step.
    n = max(wrapped_lines(v, w * 7 - 16) for v, w in zip(values, widths))
    return min(400, max(minimum, n * 16.5 + 14))


def merged(ws, area, text, size=11, bold=False, fill=WHITE, color=INK, horizontal="left"):
    ws.merge_cells(area)
    cells = ws[area]
    for rr in cells:
        for cell in rr:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = font(size, bold, color)
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
    cell = cells[0][0]
    cell.value = text
    return cell


def base_sheet(wb, title, intro, widths=BODY_WIDTHS):
    ws = wb.create_sheet(title)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.sheet_view.zoomScaleNormal = 80
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "606060"
    ws.freeze_panes = "C7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.16
    ws.page_margins.footer = 0.16
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:6"
    ws.oddHeader.right.text = f"公式情報確認 {AS_OF}"
    ws.oddHeader.right.font = FONT
    ws.oddHeader.right.size = 8
    ws.oddHeader.right.color = MID
    ws.oddFooter.left.text = "AWS M2M通信・参照設計 / 実環境の通信試験は未実施"
    ws.oddFooter.left.font = FONT
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = MID
    ws.oddFooter.right.text = "&A | &P / &N"
    ws.oddFooter.right.font = FONT
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = MID
    merged(ws, "A1:F2", title.replace("_", "  ", 1), 22, True, DARK, WHITE)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    merged(ws, "A3:F4", intro, 11, False, LIGHT)
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 23
    nav = merged(ws, "A5:B5", "← 00 読む順番", 10, False)
    nav.hyperlink = "#'00_読む順番'!A1"
    nav.font = font(10, underline="single")
    source_nav = merged(ws, "E5:F5", "一次情報一覧 → 25", 10, False, horizontal="right")
    source_nav.hyperlink = "#'25_一次情報'!A1"
    source_nav.font = font(10, underline="single")
    return ws


def header(ws, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(6, col, label)
        cell.font = font(11, True, WHITE)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[6].height = 30


def table_row(ws, r, values, widths=BODY_WIDTHS, min_height=50):
    fill = WHITE if r % 2 else "F4F4F4"
    for col, value in enumerate(values, 1):
        cell = ws.cell(r, col, value)
        cell.font = font(11 if col < 6 else 10, bold=(col == 2))
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
        cell.border = Border(bottom=Side(style="hair", color=LINE))
    ws.row_dimensions[r].height = row_height(values, widths, min_height)


def source_refs(ws, cell, refs):
    keys = re.findall(r"S\d{2}", refs)
    if not keys:
        return
    for key in keys:
        if key not in SOURCES:
            raise ValueError(f"Unknown reference {key}")
    cell.hyperlink = f"#'25_一次情報'!A{SOURCE_ROWS[keys[0]]}"
    cell.font = font(10, underline="single")
    cell.comment = Comment("\n\n".join(f"{k} {SOURCES[k]['title']}\n{SOURCES[k]['url']}" for k in keys), "出典")
    cell.comment.width = 650
    cell.comment.height = 230


def detail_sheet(wb, sec):
    ws = base_sheet(wb, sec["sheet"], sec["intro"])
    labels = ["ID", "観点・手順", "やさしい説明", "仕組み・動作原理", "設定・確認・注意", "根拠 / 出典"]
    if sec["kind"] == "sequence":
        labels = ["ID", "順番 / 接続", "通信の始点 → 終点", "行われる処理", "接続と経路の読み方", "根拠 / 出典"]
    header(ws, labels)
    for index, data in enumerate(sec["rows"], 1):
        r = index + 6
        values = [f"{sec['number']:02d}.{index:02d}", *data]
        table_row(ws, r, values)
        source_refs(ws, ws.cell(r, 6), data[4])
    ws.auto_filter.ref = f"A6:F{ws.max_row}"
    ws.print_area = f"A1:F{ws.max_row}"
    return ws


def architecture_sheet(wb, sec):
    ws = wb.create_sheet(sec["sheet"])
    for col in range(1, 21):
        ws.column_dimensions[get_column_letter(col)].width = 9.2
    for r in range(1, 63):
        ws.row_dimensions[r].height = 12
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = "A7"
    ws.sheet_properties.tabColor = DARK
    merged(ws, "A1:T4", "02  全体構成図｜認証の往復と業務の往復", 23, True, DARK, WHITE)
    nav = merged(ws, "A5:D6", "← 00 読む順番", 10)
    nav.hyperlink = "#'00_読む順番'!A1"
    nav.font = font(10, underline="single")
    merged(ws, "E5:T6", "AWS公式情報確認：2026-10-03｜同一Region・IPv4・2AZの参照設計", 11, horizontal="right")
    merged(ws, "A8:T9", "I  トークンを取得する往復（ここで一度クライアントへ戻る）", 13, True, "555555", WHITE)
    merged(ws, "A11:D15", "M2M\nクライアント\nclient ID / secret", 12, True, LIGHT, horizontal="center")
    merged(ws, "E11:F15", "⇄\nHTTPS", 11, True, horizontal="center")
    merged(ws, "G11:L15", "グローバルエッジ\nCloudFront\nCognitoが管理", 12, True, "E6E6E6", horizontal="center")
    merged(ws, "M11:N15", "⇄", 22, True, horizontal="center")
    merged(ws, "O11:T15", "Cognito User Pool\n機械を認証・scopeを確認\n署名済みJWTを発行", 12, True, LIGHT, horizontal="center")
    merged(ws, "A16:T18", "要求：POST /oauth2/token   →   応答：access token / expires_in\nCloudFrontがJWTを発行するのではなく、Cognitoが発行する。業務APIへの接続は次の別要求。", 11, fill="FAFAFA")
    merged(ws, "A20:T21", "II  取得したtokenで業務APIを呼ぶ往復", 13, True, "555555", WHITE)
    merged(ws, "A23:D27", "同じクライアント\nBearer tokenを添付\nGET /books/42", 12, True, LIGHT, horizontal="center")
    merged(ws, "E23:F27", "⇄\nHTTPS", 11, True, horizontal="center")
    merged(ws, "G23:T27", "AWSマネージドサービスの公開入口\nRegional REST API / API Gateway\nWAF → Cognito authorizer（署名・期限・scope）", 13, True, "E6E6E6", horizontal="center")
    merged(ws, "A28:T30", "公開APIへの着信は、利用者VPCのIGWを通る経路ではない。\n認可された要求をAPI GatewayがVPC Linkへ渡す。WAF/認可拒否は後段へ届かない。", 11)
    merged(ws, "A32:J33", "業務アカウントA / VPC 10.20.0.0/16", 12, True, DARK, WHITE)
    merged(ws, "M32:T33", "検査アカウントS / Security VPC", 12, True, DARK, WHITE)
    merged(ws, "A35:D40", "VPC Link V2\nENI / Link subnet\n例 10.20.10.0/24\nAPIからの出口", 11, True, LIGHT, horizontal="center")
    merged(ws, "E35:F40", "⇄\n/24\nroute", 11, True, horizontal="center")
    merged(ws, "G35:J40", "GWLBE-A\nEndpoint subnet\n10.20.20.0/24\n検査の出入口", 11, True, "DCDCDC", horizontal="center")
    merged(ws, "K35:L40", "⇄\nPrivate\nLink", 10, True, horizontal="center")
    merged(ws, "M35:T39", "GWLB + endpoint service\nフローの担当FWを選択\n検査後は元のGWLBEへ返却", 12, True, "E6E6E6", horizontal="center")
    merged(ws, "M40:T42", "⇅   GENEVE / UDP 6081\n内側：VPC Link ENI ↔ ALB の通信", 11, True, horizontal="center")
    merged(ws, "M43:T49", "Palo Alto Networks VM-Series\nセッション・security policy・脅威検査\n本文を読むにはTLS復号が必要\n戻りも同じフローを検査", 12, True, LIGHT, horizontal="center")
    merged(ws, "A41:J43", "⇅  検査済み要求をALBへ / 応答も同じGWLBEへ\n各矢印の時間順序は03章で確認", 11, True, horizontal="center")
    merged(ws, "A44:J48", "internal ALB / ALB subnet 10.20.30.0/24\nHTTPS 443を終端 → Host/pathで振り分け\nLink宛ての戻りrouteもGWLBEを指定", 12, True, "E6E6E6", horizontal="center")
    merged(ws, "A49:J51", "⇅  ALBが作る別の接続（例 HTTPS 8443）\nこの区間は今回の検査挿入の対象外", 11, True, horizontal="center")
    merged(ws, "A52:J56", "ECS task / Task subnet 10.20.40.0/24\n業務処理・データごとの権限確認\n応答はtask → ALBへ戻る", 12, True, LIGHT, horizontal="center")
    merged(ws, "M51:T56", "AZ-Aを代表表示\nAZ-Bにも同じ役割を配置\n4通りのAZ組合せの往復は\n14章のルート表・15章の動作例を参照", 11, fill="F5F5F5", horizontal="center")
    merged(ws, "A58:T62", "検査の1往復：GWLBE → GWLB → VM-Series → GWLB → 同じGWLBE\nHTTPS本文は自動で復号されない。VPC Link↔ALBとALB↔ECSは別接続。\n個別機能は公式資料に基づく。組合せ全体は設計例で、実AWS/PAN-OSの動作試験は未実施。出典：S02 / S03 / S05 / S15 / S17 / S21。", 11, fill="EAEAEA")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:T62"
    return ws


def sources_sheet(wb):
    widths = [9, 64, 62, 24, 28, 17]
    ws = base_sheet(wb, "25_一次情報", "確認日：2026-10-03。AWS・Palo Alto Networks・RFC Editorの一次資料のみ。発表時期は履歴用、現行仕様は最新Developer Guideを優先。各行のリンクから根拠を確認できます。", widths)
    header(ws, ["出典ID", "公式資料の名称", "確認した内容 / 読む際の注意", "原典へのリンク", "公開・更新", "確認日"])
    for i, (key, source) in enumerate(SOURCES.items(), 7):
        vals = [key, source["title"], source["use"], "公式ページを開く", source["published"], AS_OF]
        table_row(ws, i, vals, widths)
        cell = ws.cell(i, 4)
        cell.hyperlink = source["url"]
        cell.font = font(11, underline="single")
        cell.comment = Comment(source["url"], "URL")
    ws.auto_filter.ref = f"A6:F{ws.max_row}"
    ws.print_area = f"A1:F{ws.max_row}"
    return ws


def contents_sheet(wb):
    widths = [8, 34, 38, 78, 35, 11]
    ws = base_sheet(wb, "00_読む順番", "AWS M2M通信を、通行証と図書室のたとえから仕組み・設計・運用まで読み解く。公式情報確認日：2026-10-03。ExcelとMarkdownは同じ本文データから生成しています。", widths)
    header(ws, ["章", "クリックして移動", "読む目的", "この資料の要点", "位置付け", "項目数"])
    summaries = {
        1: "構成の可否、IGWの誤解、検査対象区間を先に把握",
        2: "認証と業務の2往復、アカウント/VPCの境界を見る",
        3: "23段階で要求と応答を順に追う",
        4: "はじめて出会うサービスをたとえで理解",
        14: "具体的なCIDRとGWLBE targetを確認",
        15: "4通りのAZ組合せで同じ検査口へ戻る理由",
        16: "HTTPSの終端と、Palo Altoで本文を読む条件",
        21: "各サービスが必要になった課題と歴史",
        23: "AWSでの受入試験の案。実機試験は未実施",
        26: "HTTP要求・応答・設定の対応を具体化",
    }
    listing = list(SECTIONS) + [dict(number=25, sheet="25_一次情報", name="一次情報", rows=list(SOURCES), intro="一次資料のリンクと、この資料で根拠とした内容。")]
    listing.sort(key=lambda s: s["number"])
    for index, sec in enumerate(listing, 7):
        num = sec["number"]
        purpose = summaries.get(num, f"{sec['name']}を詳しく確認")
        position = "図解" if num in [2, 3] else "基礎・理解" if num < 9 else "設計・動作" if num < 20 else "運用・参考"
        vals = [f"{num:02d}", sec["name"], purpose, sec["intro"], position, len(sec["rows"])]
        table_row(ws, index, vals, widths, min_height=40)
        c = ws.cell(index, 2)
        c.hyperlink = f"#'{sec['sheet']}'!A1"
        c.font = font(11, True, underline="single")
    last = ws.max_row
    notes = [
        "最初の読み順：04 たとえ → 02 全体図 → 06 認証 → 09 行き → 10 戻り → 14/15 ルート → 16 TLS。",
        "主構成：Cognito管理CloudFront → token取得。そのtokenで公開Regional REST API → WAF/認可 → VPC Link V2 → GWLBE検査 → ALB → ECS。",
        "根拠の表示：Sxxは25章の公式資料。『設計例』『推論』『未実施』は現物確認済みという意味ではありません。セルの出典リンクは最初の資料、コメントに関連URLを収録。",
        "検証の境界：文書・Excel書式・静的ルートの整合性を確認。お客様のAWS設定取得、リソース作成/変更、実通信・障害・復号試験は実施していません。",
        "Excel設定：全セルMeiryo UI／モノトーン／折り返し／先頭行と見出し固定／フィルター／出典リンク／A3横・横幅1ページ。列幅や拡大率は閲覧環境で変更できます。",
    ]
    for i, text in enumerate(notes, last + 2):
        merged(ws, f"A{i}:F{i}", text, 11, fill=LIGHT if i % 2 else WHITE)
        ws.row_dimensions[i].height = 42
    ws.auto_filter.ref = f"A6:F{last}"
    ws.print_area = f"A1:F{ws.max_row}"
    ws.freeze_panes = "C7"
    ws.sheet_properties.tabColor = DARK
    return ws


def markdown_refs(refs):
    return re.sub(r"S\d{2}", lambda m: f"[{m[0]}]({SOURCES[m[0]]['url']})", refs)


def write_markdown():
    p = [
        "# CloudFront・Cognito M2M・Regional REST API・GWLBE・Palo Alto・ALB・ECSの通信を理解する",
        "",
        f"公式情報確認日：**{AS_OF}**。AWS、Palo Alto Networks、RFC Editorの一次資料を確認した参照設計です。",
        "",
        "**最初の結論：** token取得と業務APIは別の往復です。現在のREST APIはVPC Link V2でALBへ直接統合できます。公開API Gatewayへの着信を利用者VPCのIGW経由とする説明は正確ではありません。VPC LinkとALBの間へGWLBEを挿入するには、別subnet・往復の具体的なルート・FWの対称性が必要です。",
        "",
        markdown_refs("根拠：S01 S02 S13 S14 S15 S19。"),
        "",
        "前提は東京リージョン、IPv4、2AZ、同一業務VPC、internal ALB、ECS awsvpcです。住所・IP・portは架空の設計例です。実際のAWS設定は取得しておらず、実通信、TLS復号、性能、障害の試験は未実施です。個別機能の確認と、組合せの設計上の推論を区別します。",
        "",
        "推奨する読み順：04のたとえ → 02の全体図 → 06の認証 → 09/10の往復 → 14/15のルート → 16のTLS。Excelは全セルMeiryo UIとモノトーンで、同じ本文を章別のシートに収録しています。",
        "",
        "## 目次", "",
    ]
    for sec in SECTIONS:
        if sec["number"] == 26:
            p.append("- [25 一次情報](#chapter-25)")
        p.append(f"- [{sec['number']:02d} {sec['name']}](#chapter-{sec['number']:02d})")
    theme = {"theme": "base", "themeVariables": {
        "fontFamily": FONT, "primaryColor": "#eeeeee", "primaryTextColor": "#202020", "primaryBorderColor": "#555555",
        "lineColor": "#555555", "secondaryColor": "#f5f5f5", "tertiaryColor": "#ffffff", "clusterBkg": "#f6f6f6",
        "clusterBorder": "#777777", "actorBkg": "#eeeeee", "actorBorder": "#555555", "actorTextColor": "#202020",
        "actorLineColor": "#777777", "signalColor": "#444444", "signalTextColor": "#202020", "noteBkgColor": "#f2f2f2",
        "noteBorderColor": "#777777", "noteTextColor": "#202020", "labelBoxBkgColor": "#eeeeee", "labelTextColor": "#202020",
    }}
    for sec in SECTIONS:
        if sec["number"] == 26:
            append_sources(p)
        p.extend(["", f'<a id="chapter-{sec["number"]:02d}"></a>', "", f"## {sec['number']:02d} {sec['name']}", "", sec["intro"], ""])
        if sec["kind"] in ["architecture", "sequence"]:
            diagram = content_diagrams.ARCHITECTURE if sec["kind"] == "architecture" else content_diagrams.SEQUENCE
            p.extend(["```mermaid", "%%{init: " + json.dumps(theme, ensure_ascii=False) + "}%%", diagram.replace(r"\n", "<br/>").strip(), "```", ""])
            if sec["kind"] == "architecture":
                p.extend(["Mermaid非対応の表示環境でも、次の文字図と03章の順序表で経路を確認できます。", "", "```text", "認証：client ⇄ CloudFront ⇄ Cognito（tokenはclientへ戻る）", "業務：client ⇄ Regional REST API［WAF → Cognito認可］ ⇄ VPC Link ENI", "      ⇄ GWLBE ⇄ internal ALB ⇄ ECS task", "検査：GWLBE → GWLB → VM-Series → GWLB → 同じGWLBE", "      業務要求と業務応答の両方向に上記検査を適用。", "```", ""])
        for i, (topic, easy, detail, practical, refs) in enumerate(sec["rows"], 1):
            p.extend([f"### {sec['number']:02d}.{i:02d} {topic}", "", f"**やさしく言うと：** {easy}", ""])
            if sec["number"] == 26:
                p.extend(["```text", detail, "```", ""])
            else:
                p.extend([detail.replace("\n", "  \n"), ""])
            p.extend([f"**設定・確認：** {practical}", "", markdown_refs(f"根拠・位置付け：{refs}"), ""])
    p.extend(["---", "", "この資料は公式仕様をもとにした解説と参照設計です。実環境への適用は、23章の確認を実際のバージョン・設定で行って判断します。資料作成時にAWSリソースの変更や秘密情報の取得はしていません。", ""])
    MD.write_text("\n".join(p), encoding="utf-8", newline="\n")


def append_sources(p):
    p.extend(["", '<a id="chapter-25"></a>', "", "## 25 一次情報", "", f"全資料の確認日：{AS_OF}。継続更新文書は将来変更されるため、適用時にも再確認してください。古い発表記事は歴史の根拠として使用しています。", ""])
    for key, s in SOURCES.items():
        p.extend([f"### {key} {s['title']}", "", f"[公式資料を開く]({s['url']})  ", f"公開・更新：{s['published']}  ", f"確認した内容：{s['use']}", ""])


def main():
    wb = Workbook()
    wb.remove(wb.active)
    wb._fonts[0] = font()
    wb._named_styles[0].font = font()
    wb.properties.title = "AWS M2M通信・GWLB検査経路 詳細解説"
    wb.properties.subject = "CloudFront / Cognito / API Gateway REST / VPC Link V2 / GWLBE / VM-Series / ALB / ECS"
    wb.properties.creator = "Architecture Guide"
    wb.properties.description = "2026-10-03公式情報確認。日本語の参照設計。実AWS通信試験は未実施。"
    wb.properties.created = datetime(2026, 10, 3)
    wb.properties.modified = datetime(2026, 10, 3)
    contents_sheet(wb)
    for sec in SECTIONS:
        if sec["number"] == 26:
            sources_sheet(wb)
        if sec["kind"] == "architecture":
            architecture_sheet(wb, sec)
        else:
            detail_sheet(wb, sec)
    # Keep the Office theme grayscale as well as all explicitly styled cells.
    from openpyxl.writer.theme import theme_xml
    gray = iter(["404040", "606060", "808080", "A0A0A0", "C0C0C0", "E0E0E0"])
    themed = re.sub(r"(<a:accent[1-6]>\s*<a:srgbClr val=\")[A-Fa-f0-9]+", lambda m: m[1] + next(gray), theme_xml)
    themed = themed.replace('val="0000FF"', 'val="303030"').replace('val="800080"', 'val="606060"')
    wb.loaded_theme = themed
    wb.active = 0
    wb.save(XLSX)
    write_markdown()
    print(json.dumps({"xlsx": str(XLSX), "markdown": str(MD), "sheets": len(wb.sheetnames), "entries": sum(len(s["rows"]) for s in SECTIONS), "sources": len(SOURCES)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
